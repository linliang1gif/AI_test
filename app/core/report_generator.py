"""Generate a Word test case report document."""
from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Dict, List

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT


def generate_report(
    cases: List[dict],
    output_path: Path,
    project_name: str = "",
    module_name: str = "",
    coverage: Dict = None,
) -> Path:
    """Generate a professional Word test case report."""
    doc = Document()

    # Title
    title = doc.add_heading("测试用例报告", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Meta info
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(f"项目：{project_name or '未指定'}").font.size = Pt(12)
    meta.add_run(f"    模块：{module_name or '未指定'}").font.size = Pt(12)
    meta.add_run(f"    日期：{datetime.now().strftime('%Y-%m-%d')}").font.size = Pt(12)

    doc.add_paragraph("")

    # Coverage summary
    if coverage:
        doc.add_heading("覆盖率概览", level=1)
        qs = coverage.get("quality_score", 0)
        doc.add_paragraph(f"质量评分：{qs}/100")
        doc.add_paragraph(f"用例总数：{coverage.get('total', len(cases))}")

        dims = coverage.get("dimensions", {})
        if dims:
            doc.add_paragraph("覆盖维度分布：")
            for dim, count in dims.items():
                doc.add_paragraph(f"  • {dim}：{count} 条", style="List Bullet")

        missing = coverage.get("missing_dimensions", [])
        if missing:
            p = doc.add_paragraph("未覆盖维度：")
            p.add_run("、".join(missing)).font.color.rgb = RGBColor(0xFF, 0x00, 0x00)

        pri = coverage.get("priority_dist", {})
        if pri:
            doc.add_paragraph(f"优先级分布：高({pri.get('高', 0)}) / 中({pri.get('中', 0)}) / 低({pri.get('低', 0)})")

    # Test cases table
    doc.add_heading("测试用例明细", level=1)

    headers = ["序号", "测试点", "用例标题", "前置条件", "操作步骤", "预期结果", "优先级"]
    col_widths = [Inches(0.5), Inches(1.2), Inches(1.5), Inches(1.2), Inches(1.8), Inches(1.8), Inches(0.5)]

    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Header row
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.font.bold = True
                run.font.size = Pt(9)

    # Data rows
    for idx, case in enumerate(cases, 1):
        row = table.add_row()
        values = [
            str(idx),
            case.get("测试点", case.get("功能点", "")),  # 支持新旧格式
            case.get("用例标题", ""),
            case.get("前置条件", ""),
            case.get("操作步骤", case.get("测试步骤", "")),
            case.get("预期结果", ""),
            case.get("优先级", "中"),
        ]
        for i, val in enumerate(values):
            cell = row.cells[i]
            cell.text = val
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(8)

    # Set column widths
    for i, width in enumerate(col_widths):
        for row in table.rows:
            row.cells[i].width = width

    doc.add_paragraph("")
    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run(f"报告生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}").font.size = Pt(8)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return output_path
