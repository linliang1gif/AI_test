from __future__ import annotations

import zipfile
from pathlib import Path
from typing import List, Optional

from docx import Document

from app.core.ocr import OcrService

# Supported file extensions
SUPPORTED_EXTENSIONS = {".docx", ".pdf", ".txt", ".md"}


def is_url(text: str) -> bool:
    """Check if a string looks like a URL."""
    return text.strip().startswith(("http://", "https://"))


def read_document(file_path: Path, ocr_service: Optional[OcrService] = None) -> str:
    """Read any supported document format and return text content."""
    suffix = file_path.suffix.lower()
    if suffix == ".docx":
        return read_docx(file_path, ocr_service=ocr_service)
    elif suffix == ".pdf":
        return read_pdf(file_path)
    elif suffix in (".txt", ".md"):
        return read_text(file_path)
    else:
        raise ValueError(f"不支持的文件格式：{suffix}")


def read_url(
    url: str,
    username: Optional[str] = None,
    password: Optional[str] = None,
) -> str:
    """Fetch a web page and return its text content."""
    from app.core.web_loader import fetch_page
    page = fetch_page(url, username=username, password=password)
    return page.text


def read_docx(file_path: Path, ocr_service: Optional[OcrService] = None) -> str:
    """Return concatenated text from a Word document, with optional OCR."""
    document = Document(file_path)

    # Extract text from paragraphs
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]

    # Also extract text from tables
    table_texts = []
    for table in document.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                table_texts.append(row_text)

    text_parts = ["\n".join(paragraphs)]
    if table_texts:
        text_parts.append("[表格内容]\n" + "\n".join(table_texts))

    if ocr_service:
        image_texts = _extract_images_text(file_path, ocr_service)
        if image_texts:
            text_parts.append("\n".join(image_texts))

    return "\n\n".join([part for part in text_parts if part])


def read_pdf(file_path: Path) -> str:
    """Read text from a PDF file."""
    try:
        import pdfplumber
    except ImportError:
        raise ImportError(
            "读取 PDF 需要安装 pdfplumber：pip install pdfplumber"
        )

    texts: List[str] = []
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if text:
                texts.append(text.strip())
            # Also extract tables
            tables = page.extract_tables()
            for table in tables:
                for row in table:
                    row_text = " | ".join(str(cell or "").strip() for cell in row)
                    if row_text.strip():
                        texts.append(row_text)
    return "\n\n".join(texts)


def read_text(file_path: Path) -> str:
    """Read plain text or markdown file."""
    encodings = ["utf-8", "gbk", "gb2312", "utf-16", "latin-1"]
    for enc in encodings:
        try:
            return file_path.read_text(encoding=enc).strip()
        except (UnicodeDecodeError, UnicodeError):
            continue
    return file_path.read_text(encoding="utf-8", errors="replace").strip()


def _extract_images_text(file_path: Path, ocr_service: OcrService) -> List[str]:
    texts: List[str] = []
    with zipfile.ZipFile(file_path) as docx_zip:
        image_files = [f for f in docx_zip.namelist() if f.startswith("word/media/")]
        for index, image_name in enumerate(image_files, 1):
            image_bytes = docx_zip.read(image_name)
            ocr_text = ocr_service.recognize(image_bytes)
            if ocr_text:
                texts.append(f"[图片{index}] {ocr_text}")
    return texts


def get_file_filter() -> str:
    """Return file dialog filter string for supported formats."""
    return (
        "所有支持格式 (*.docx *.pdf *.txt *.md);;"
        "Word 文档 (*.docx);;"
        "PDF 文档 (*.pdf);;"
        "文本文件 (*.txt *.md)"
    )


def is_supported(path: Path) -> bool:
    """Check if a file is a supported document format."""
    return path.suffix.lower() in SUPPORTED_EXTENSIONS
