#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
直接测试API接口,获取详细的错误信息
"""

import requests
import json

def test_api():
    """测试API接口"""
    
    print("=" * 80)
    print("API接口测试 - 详细诊断")
    print("=" * 80)
    
    # 测试1: 检查服务是否运行
    print("\n1. 检查后端服务状态...")
    try:
        response = requests.get("http://127.0.0.1:8081/api/dashboard/stats", timeout=5)
        print(f"   ✅ 后端服务运行正常 (状态码: {response.status_code})")
    except Exception as e:
        print(f"   ❌ 后端服务异常: {e}")
        return
    
    # 测试2: 测试文本文件上传
    print("\n2. 测试文本文件上传...")
    test_content = """
电商平台需求文档

1. 用户管理
   - 用户登录
   - 用户注册
   - 密码找回

2. 商品管理
   - 商品搜索
   - 商品详情
   - 商品分类

3. 购物车
   - 添加商品
   - 修改数量
   - 删除商品

4. 订单管理
   - 创建订单
   - 订单支付
   - 订单查询
    """
    
    files = {'file': ('需求文档.txt', test_content.encode('utf-8'), 'text/plain')}
    
    try:
        response = requests.post(
            'http://127.0.0.1:8081/api/testcases/generate',
            files=files,
            timeout=30
        )
        
        print(f"   状态码: {response.status_code}")
        print(f"   响应头: {dict(response.headers)}")
        
        try:
            result = response.json()
            print(f"   响应内容: {json.dumps(result, ensure_ascii=False, indent=2)}")
            
            if result.get('success'):
                print(f"   ✅ 测试通过 - 生成了 {result.get('count', 0)} 个测试用例")
            else:
                print(f"   ❌ 测试失败 - 错误: {result.get('error', '未知错误')}")
        except:
            print(f"   响应文本: {response.text}")
            
    except Exception as e:
        print(f"   ❌ 请求异常: {e}")
    
    # 测试3: 测试Word文档上传(如果有docx库)
    print("\n3. 测试Word文档上传...")
    try:
        import docx
        
        # 创建Word文档
        doc = docx.Document()
        doc.add_heading('测试需求', 0)
        doc.add_paragraph('用户登录功能')
        doc.add_paragraph('商品搜索功能')
        doc.add_paragraph('订单支付功能')
        
        doc_path = 'test_doc.docx'
        doc.save(doc_path)
        
        with open(doc_path, 'rb') as f:
            files = {'file': (doc_path, f, 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')}
            
            response = requests.post(
                'http://127.0.0.1:8081/api/testcases/generate',
                files=files,
                timeout=30
            )
            
            print(f"   状态码: {response.status_code}")
            
            try:
                result = response.json()
                if result.get('success'):
                    print(f"   ✅ Word文档测试通过 - 生成了 {result.get('count', 0)} 个测试用例")
                else:
                    print(f"   ❌ Word文档测试失败 - 错误: {result.get('error', '未知错误')}")
            except:
                print(f"   响应文本: {response.text}")
                
    except ImportError:
        print("   ⚠️  python-docx未安装,跳过Word文档测试")
    except Exception as e:
        print(f"   ❌ Word文档测试异常: {e}")
    
    # 测试4: 检查后端日志
    print("\n4. 建议检查后端日志:")
    print("   - 查看是否有文档解析错误")
    print("   - 查看是否有依赖库缺失警告")
    print("   - 查看是否有其他异常信息")
    
    print("\n" + "=" * 80)
    print("测试完成")
    print("=" * 80)


if __name__ == "__main__":
    test_api()
