#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
端到端测试 - 模拟前端完整上传流程
"""

import requests
import json
import time

def test_e2e():
    """端到端测试"""
    
    print("\n" + "=" * 80)
    print("端到端测试 - 模拟前端完整流程")
    print("=" * 80 + "\n")
    
    # 步骤1: 检查后端服务
    print("步骤1: 检查后端服务...")
    try:
        response = requests.get("http://127.0.0.1:8081/api/dashboard/stats", timeout=5)
        if response.status_code == 200:
            print("✅ 后端服务正常\n")
        else:
            print(f"❌ 后端服务异常: {response.status_code}\n")
            return
    except Exception as e:
        print(f"❌ 无法连接后端服务: {e}\n")
        return
    
    # 步骤2: 准备测试文件
    print("步骤2: 准备测试文件...")
    test_files = []
    
    # 2.1 文本文件
    txt_content = """
电商平台需求文档

功能模块:
1. 用户登录
2. 商品搜索  
3. 购物车管理
4. 订单支付
    """
    test_files.append(('需求.txt', txt_content.encode('utf-8'), 'text/plain'))
    print("✅ 准备文本文件")
    
    # 2.2 Word文档(如果可用)
    try:
        import docx
        doc = docx.Document()
        doc.add_heading('需求文档', 0)
        doc.add_paragraph('用户登录功能')
        doc.add_paragraph('商品管理功能')
        doc_path = 'test_需求.docx'
        doc.save(doc_path)
        with open(doc_path, 'rb') as f:
            test_files.append((doc_path, f.read(), 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'))
        print("✅ 准备Word文档")
    except:
        print("⚠️  跳过Word文档(python-docx未安装)")
    
    print()
    
    # 步骤3: 测试每个文件
    for i, (filename, content, content_type) in enumerate(test_files, 1):
        print(f"步骤3.{i}: 测试上传 {filename}...")
        
        files = {'file': (filename, content, content_type)}
        
        try:
            # 发送请求
            start_time = time.time()
            response = requests.post(
                'http://127.0.0.1:8081/api/testcases/generate',
                files=files,
                timeout=30
            )
            elapsed = time.time() - start_time
            
            print(f"  - 状态码: {response.status_code}")
            print(f"  - 响应时间: {elapsed:.2f}秒")
            print(f"  - Content-Type: {response.headers.get('content-type', 'N/A')}")
            
            # 解析响应
            if response.status_code == 200:
                try:
                    result = response.json()
                    
                    if result.get('success'):
                        print(f"  ✅ 成功生成 {result.get('count', 0)} 个测试用例")
                        
                        # 显示部分测试用例
                        if 'testCases' in result and result['testCases']:
                            print(f"  - 示例测试用例:")
                            for tc in result['testCases'][:2]:
                                print(f"    • {tc.get('title', 'N/A')}")
                    else:
                        print(f"  ❌ 生成失败: {result.get('error', '未知错误')}")
                        print(f"  - 完整响应: {json.dumps(result, ensure_ascii=False, indent=2)}")
                        
                except json.JSONDecodeError:
                    print(f"  ❌ 响应不是有效的JSON")
                    print(f"  - 响应内容: {response.text[:200]}")
            else:
                print(f"  ❌ HTTP错误: {response.status_code}")
                print(f"  - 响应内容: {response.text[:200]}")
                
        except requests.Timeout:
            print(f"  ❌ 请求超时")
        except Exception as e:
            print(f"  ❌ 请求异常: {e}")
        
        print()
    
    # 步骤4: 验证生成的测试用例
    print("步骤4: 验证生成的测试用例...")
    try:
        response = requests.get("http://127.0.0.1:8081/api/test-cases", timeout=5)
        if response.status_code == 200:
            result = response.json()
            test_cases = result.get('testCases', [])
            ai_generated = [tc for tc in test_cases if 'AI生成' in tc.get('title', '')]
            print(f"✅ 当前系统中有 {len(test_cases)} 个测试用例")
            print(f"✅ 其中 {len(ai_generated)} 个是AI生成的")
        else:
            print(f"⚠️  无法获取测试用例列表: {response.status_code}")
    except Exception as e:
        print(f"⚠️  获取测试用例列表异常: {e}")
    
    print("\n" + "=" * 80)
    print("测试完成")
    print("=" * 80 + "\n")
    
    # 步骤5: 给出诊断建议
    print("诊断建议:")
    print("1. 如果后端测试全部通过,但前端仍然报错,请:")
    print("   - 打开浏览器开发者工具(F12)")
    print("   - 查看Network标签页")
    print("   - 重新上传文件")
    print("   - 查看请求的详细信息(Request/Response)")
    print()
    print("2. 常见问题:")
    print("   - CORS跨域问题: 检查后端CORS配置")
    print("   - 文件大小限制: 检查是否超过上传限制")
    print("   - 文件格式问题: 确认文件格式正确")
    print("   - 网络问题: 检查前后端是否在同一网络")
    print()


if __name__ == "__main__":
    test_e2e()
