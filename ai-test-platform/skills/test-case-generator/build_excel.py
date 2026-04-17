#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
v1.2.2 完整测试用例汇总导出
运行: py build_excel.py
生成: v1.2.2_测试用例_complete.xlsx
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from collections import Counter
import os, sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import case_helper
import cases_m1
import cases_m2
import cases_m3
import cases_m5
import cases_m6m7
import cases_m8m9
import cases_m10m11
import cases_supplement
import cases_security
import cases_interface
import cases_perf
import cases_compat
import cases_m1_extra
import cases_m5_extra

ROWS = case_helper.ROWS

HDR = ['序号','模块名称','测试点','用例标题','优先级','前置条件','测试步骤','测试数据','预期结果','测试类型']
COL_W = [6, 22, 28, 38, 8, 40, 50, 40, 40, 12]

MOD_COLORS = {
    '付款单-新增字段':       'DDEEFF',
    '付款单-其他优化':       'E8F5E9',
    '付款单-支付设置联动':   'FFF9C4',
    '付款单-实付退总金额计算':'FCE4EC',
    '付款单-分账对接招行银联':'F3E5F5',
    '企业管理-系统管理入口':  'E0F7FA',
    '企业管理-付款单支付设置':'FFF3E0',
    '企业管理-反向开票设置':  'E8EAF6',
    '应付单-新增不含税应付金额':'F1F8E9',
    '供应商管理-个税承担方':  'FBE9E7',
    '采购订单-个税承担方':   'E3F2FD',
    '反向开票申请单-先付款后开票联动':'EDE7F6',
    '付款单-安全测试':            'FFEBEE',
    '付款单-接口测试':            'E8F5E9',
    '付款单-性能测试':            'FFF8E1',
    '付款单-兼容性测试':          'F3E5F5',
}
PRI_COLORS = {'高':'FF5252','中':'FF9800','低':'4CAF50','边界':'2196F3'}

thin = Side(style='thin', color='BBBBBB')
border = Border(left=thin, right=thin, top=thin, bottom=thin)

def make_excel(rows, out_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'v1.2.2测试用例'
    ws.freeze_panes = 'A2'

    for ci, h in enumerate(HDR, 1):
        cell = ws.cell(1, ci, h)
        cell.font = Font(bold=True, color='FFFFFF', size=11)
        cell.fill = PatternFill('solid', fgColor='1565C0')
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = border
    ws.row_dimensions[1].height = 28

    for ci, w in enumerate(COL_W, 1):
        ws.column_dimensions[get_column_letter(ci)].width = w

    for row in rows:
        ri = row[0] + 1
        mod = row[1]
        bg = MOD_COLORS.get(mod, 'FFFFFF')
        for ci, val in enumerate(row, 1):
            cell = ws.cell(ri, ci, val)
            cell.fill = PatternFill('solid', fgColor=bg)
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            cell.border = border
            if ci == 1:
                cell.alignment = Alignment(horizontal='center', vertical='top')
            if ci == 5:
                color = PRI_COLORS.get(str(val), '000000')
                cell.font = Font(bold=True, color=color)
        ws.row_dimensions[ri].height = 80

    ws2 = wb.create_sheet('统计')
    ws2.column_dimensions['A'].width = 38
    ws2.column_dimensions['B'].width = 10
    ws2.cell(1,1,'模块').font = Font(bold=True)
    ws2.cell(1,2,'用例数').font = Font(bold=True)
    cnt = Counter(r[1] for r in rows)
    for i, (mod, n) in enumerate(cnt.items(), 2):
        ws2.cell(i, 1, mod)
        ws2.cell(i, 2, n)
    total_row = len(cnt) + 2
    ws2.cell(total_row, 1, '合计').font = Font(bold=True)
    ws2.cell(total_row, 2, len(rows)).font = Font(bold=True)

    wb.save(out_path)
    print(f'导出完成：{len(rows)} 条用例 -> {out_path}')
    # 打印各模块统计
    print('\n各模块用例数：')
    for mod, n in cnt.items():
        print(f'  {mod}: {n}')

if __name__ == '__main__':
    out = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'v1.2.2_测试用例_complete.xlsx')
    make_excel(ROWS, out)
