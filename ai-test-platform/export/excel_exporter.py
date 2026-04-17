#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
AI Test Platform - Excel导出器

将测试用例导出为Excel格式。
"""

from pathlib import Path
from typing import List, Dict, Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from config.config import get_config

class ExcelExporter:
    """Excel导出器"""
    
    def __init__(self):
        self.config = get_config()
    
    def export_testcases(self, testcases: List[Dict[str, Any]], output_file: str = None) -> Path:
        """导出测试用例到Excel"""
        if output_file is None:
            output_file = self.config.paths.output_dir / "testcases.xlsx"
        
        output_path = Path(output_file)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        
        # 创建工作簿
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "测试用例"
        
        # 设置表头
        headers = [
            "序号", "模块名称", "测试点", "用例标题", "前置条件", 
            "测试步骤", "测试数据", "预期结果", "优先级", "测试类型",
            "复杂度", "预估时间", "自动化可行性", "风险等级", "标签"
        ]
        
        # 写入表头
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )
        
        # 写入数据
        for row, testcase in enumerate(testcases, 2):
            data = [
                testcase.get('序号', ''),
                testcase.get('模块名称', ''),
                testcase.get('测试点', ''),
                testcase.get('用例标题', ''),
                testcase.get('前置条件', ''),
                testcase.get('测试步骤', ''),
                testcase.get('测试数据', ''),
                testcase.get('预期结果', ''),
                testcase.get('优先级', ''),
                testcase.get('测试类型', ''),
                testcase.get('复杂度', ''),
                testcase.get('预估时间', ''),
                testcase.get('自动化可行性', ''),
                testcase.get('风险等级', ''),
                testcase.get('标签', '')
            ]
            
            for col, value in enumerate(data, 1):
                cell = ws.cell(row=row, column=col, value=str(value))
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                cell.border = Border(
                    left=Side(style='thin'),
                    right=Side(style='thin'),
                    top=Side(style='thin'),
                    bottom=Side(style='thin')
                )
                
                # 根据优先级设置颜色
                if col == 9:  # 优先级列
                    if value == '高':
                        cell.fill = PatternFill(start_color="FFE6E6", end_color="FFE6E6", fill_type="solid")
                    elif value == '中':
                        cell.fill = PatternFill(start_color="FFF2E6", end_color="FFF2E6", fill_type="solid")
        
        # 调整列宽
        column_widths = [8, 15, 20, 30, 25, 40, 20, 30, 10, 12, 10, 12, 15, 10, 20]
        for col, width in enumerate(column_widths, 1):
            ws.column_dimensions[get_column_letter(col)].width = width
        
        # 设置行高
        for row in range(2, len(testcases) + 2):
            ws.row_dimensions[row].height = 60
        
        # 冻结首行
        ws.freeze_panes = "A2"
        
        # 保存文件
        wb.save(output_path)
        
        print(f"📊 测试用例已导出到Excel: {output_path}")
        return output_path