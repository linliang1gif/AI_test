#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
测试上传Word文档生成测试用例
"""

import requests

def test_docx_upload():
    """测试Word文档上传"""
    
    print("🧪 测试Word文档上传\n")
    
    # 检查是否有docx库
    try:
        import docx
        print("✅ python-docx已安装")
        
        # 创建一个Word文档
        doc = docx.Document()
        doc.add_heading('电商平台需求文档', 0)
        
        doc.add_heading('1. 用户管理模块', level=1)
        doc.add_paragraph('用户注册功能')
        doc.add_paragraph('用户登录功能')
        doc.add_paragraph('密码找回功能')
        
        doc.add_heading('2. 商品管理模块', level=1)
        doc.add_paragraph('商品搜索功能')
        doc.add_paragraph('商品详情查看')
        doc.add_paragraph('商品分类浏览')
        
        doc.add_heading('3. 购物车模块', level=1)
        doc.add_paragraph('添加商品到购物车')
        doc.add_paragraph('修改购物车商品数量')
        doc.add_paragraph('删除购物车商品')
        
        # 保存文档
        doc_path = 'test_requirement.docx'
        doc.save(doc_path)
        print(f"✅ 创建测试文档: {doc_path}\n")
        
        # 上传文档
        with open(doc_path, 'rb') as f:
            files = {'file': (doc_path, f, 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')}
            
            print("📤 上传Word文档...")
            response = requests.post(
                'http://127.0.0.1:8081/api/testcases/generate',
                files=files,
                timeout=30
            )
            
            print(f"状态码: {response.status_code}")
            
            if response.status_code == 200:
                result = response.json()
                print("\n✅ 测试用例生成成功!")
                print(f"生成数量: {result.get('count', 0)}")
                print(f"消息: {result.get('message', '')}")
                
                if 'testCases' in result:
                    print(f"\n生成的测试用例:")
                    for i, tc in enumerate(result['testCases'][:5], 1):
                        print(f"\n{i}. {tc.get('title', '')}")
                        print(f"   优先级: {tc.get('priority', '')}")
                        print(f"   模块: {tc.get('module', '')}")
            else:
                print(f"\n❌ 请求失败")
                print(f"响应: {response.text}")
                
    except ImportError:
        print("❌ python-docx未安装")
        print("请运行: pip install python-docx")
    except Exception as e:
        print(f"\n❌ 错误: {e}")


if __name__ == "__main__":
    test_docx_upload()
